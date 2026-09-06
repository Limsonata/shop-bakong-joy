#!/usr/bin/env python3
"""
One-off importer: Clothing_Shop_Management3.xlsx  ->  POS seed data.

Reads the shop's Excel workbook and produces:
  * src/data/shop-seed.json                       (used by the local/demo backend)
  * supabase/migrations/<ts>_pos_seed.sql         (used when Supabase is connected)

Usage:
    python3 -m venv .venv && .venv/bin/pip install openpyxl
    .venv/bin/python tools/import_excel.py ~/Downloads/Clothing_Shop_Management3.xlsx

Sheets consumed:
    Stock              -> stock_items
    Order2             -> sales + sale_items   (multi-row orders grouped by Order ID)
    Orders             -> address / city / delivery method for those sales
    Income & Expenses  -> finance_entries
    Profit             -> ignored (recomputed by the app)
"""

from __future__ import annotations

import json
import re
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required:  pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
JSON_OUT = ROOT / "src" / "data" / "shop-seed.json"
SQL_OUT = ROOT / "supabase" / "migrations" / "20260906090100_pos_seed.sql"

# Rows that are leftover template/sample data in the workbook, not real records.
JUNK_MARKERS = ("basic t-shirt", "sample customer", "sample")


# ─────────────────────────────── helpers ────────────────────────────────
def clean(value) -> str:
    """Trim, collapse whitespace, strip zero-width characters."""
    if value is None:
        return ""
    text = str(value).replace("\u200b", " ")
    text = unicodedata.normalize("NFC", text)
    return re.sub(r"\s+", " ", text).strip()


def num(value, default=0.0) -> float:
    if value is None or value == "":
        return default
    if isinstance(value, (int, float)):
        return round(float(value), 4)
    try:
        return round(float(str(value).replace(",", "").strip()), 4)
    except ValueError:
        return default


def as_int(value, default=0) -> int:
    return int(round(num(value, default)))


def iso_date(value, fallback: str) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = clean(value)
    match = re.match(r"(\d{4})-(\d{2})-(\d{2})", text)
    return f"{match.group(1)}-{match.group(2)}-{match.group(3)}" if match else fallback


def looks_like_date(value) -> bool:
    return isinstance(value, (datetime, date)) or bool(re.match(r"^\d{4}-\d{2}-\d{2}", clean(value)))


def is_junk(*values) -> bool:
    blob = " ".join(clean(v).lower() for v in values)
    return any(marker in blob for marker in JUNK_MARKERS)


def norm_size(value) -> str:
    text = clean(value)
    if not text:
        return "One size"
    if "free" in text.lower():
        return "Free size"
    return text.upper() if len(text) <= 2 and text.isalpha() else text


def norm_color(value) -> str:
    text = clean(value)
    return text[:1].upper() + text[1:] if text else ""


CATEGORY_RULES = [
    ("Underwear", ("underwear", "underwar", "under wear")),
    ("Bra & Lingerie", ("bra", "lingerie", "ligerie", "lace", "satin", "sbai", "desir")),
    ("Jeans", ("jean", "denim")),
    ("Tops", ("top", "shirt", "tank", "zara")),
]


def guess_category(name: str) -> str:
    low = name.lower()
    for category, keywords in CATEGORY_RULES:
        if any(keyword in low for keyword in keywords):
            return category
    return "Other"


def match_key(name: str, size: str) -> str:
    """Loose key used to link an order line back to a stock SKU."""
    base = re.sub(r"[^a-z0-9]", "", clean(name).lower())
    for noise in ("ligerie", "lingerie"):
        base = base.replace(noise, "lace")
    return f"{base}|{norm_size(size).lower()}"


# ──────────────────────────────── Stock ─────────────────────────────────
def read_stock(wb) -> list[dict]:
    ws = wb["Stock"]
    items: list[dict] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        sku, name, _sheet_category, size, color, cost, price, stock_in, sold = row[:9]
        name = clean(name)
        if not clean(sku) or not name or is_junk(sku, name):
            continue
        items.append(
            {
                "sku": clean(sku),
                "name": name,
                "category": guess_category(name),
                "size": norm_size(size),
                "color": norm_color(color),
                "cost": num(cost),
                "price": num(price),
                "stockIn": as_int(stock_in),
                "sold": as_int(sold),
            }
        )
    return items


# ────────────────────── Orders sheet (delivery details) ─────────────────
def read_delivery_details(wb) -> dict[str, dict]:
    """customer-name -> {address, city, method} taken from the wider `Orders` sheet."""
    ws = wb["Orders"]
    details: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        customer, phone, address, city = row[2], row[3], row[4], row[5]
        method = row[16]
        name = clean(customer)
        if not name or is_junk(customer, address, city):
            continue
        details[name.lower()] = {
            "address": clean(address),
            "city": clean(city).title(),
            "method": clean(method) or "Delivery",
            "phone": clean(phone),
        }
    return details


# ──────────────────────────── Order2 -> sales ───────────────────────────
DELIVERY_STATUS = {"preparing": "preparing", "done": "delivered", "paid": "delivered"}
ORDER_STATUS = {"confirmed": "confirmed", "cancelled": "cancelled"}


def read_sales(wb, stock_by_key: dict[str, dict], delivery: dict[str, dict]) -> list[dict]:
    ws = wb["Order2"]
    groups: dict[str, dict] = {}
    order_sequence: list[str] = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        (
            code,
            row_date,
            customer,
            phone,
            product,
            size,
            qty,
            price,
            _total,
            deposit,
            _balance,
            delivery_fee,
            delivery_status,
            order_status,
            note,
        ) = row[:15]

        code = clean(code)
        product = clean(product)
        if not code or not product or is_junk(code, customer, product):
            continue

        if code not in groups:
            order_sequence.append(code)
            groups[code] = {
                "code": code,
                "date": None,
                "customer": "",
                "phone": "",
                "deposit": 0.0,
                "deliveryFee": 0.0,
                "deliveryStatus": "preparing",
                "status": "open",
                "note": "",
                "items": [],
            }
        group = groups[code]

        # Header values only appear on the first line of an order.
        if row_date is not None and group["date"] is None:
            group["date"] = iso_date(row_date, "")
        raw_customer = clean(customer)
        if raw_customer and not group["customer"]:
            # One row has a date typed into the customer column — keep it out.
            group["customer"] = "" if looks_like_date(customer) else raw_customer
        if clean(phone) and not group["phone"]:
            group["phone"] = clean(phone).replace(" ", "")
        if deposit is not None:
            group["deposit"] = max(group["deposit"], num(deposit))
        if delivery_fee is not None:
            group["deliveryFee"] = max(group["deliveryFee"], num(delivery_fee))
        if clean(delivery_status):
            key = clean(delivery_status).lower()
            group["deliveryStatus"] = DELIVERY_STATUS.get(key, "preparing")
        if clean(order_status):
            group["status"] = ORDER_STATUS.get(clean(order_status).lower(), "confirmed")
        if clean(note):
            group["note"] = clean(note)

        match = stock_by_key.get(match_key(product, size))
        group["items"].append(
            {
                "sku": match["sku"] if match else None,
                "name": product,
                "size": norm_size(size),
                "color": match["color"] if match else "",
                "qty": max(as_int(qty, 1), 1),
                "unitPrice": num(price),
                "unitCost": match["cost"] if match else 0.0,
            }
        )

    sales: list[dict] = []
    for index, code in enumerate(order_sequence, start=1):
        group = groups[code]
        subtotal = round(sum(item["qty"] * item["unitPrice"] for item in group["items"]), 2)
        total = round(subtotal + group["deliveryFee"], 2)
        paid = min(group["deposit"], total) if group["deposit"] else 0.0
        contact = delivery.get(group["customer"].lower(), {})

        sales.append(
            {
                "code": f"S-{index:04d}",
                "legacyCode": code,
                "date": group["date"] or "2026-09-01",
                "customerName": group["customer"] or "Walk-in customer",
                "phone": group["phone"] or contact.get("phone", ""),
                "address": contact.get("address", ""),
                "city": contact.get("city", ""),
                "channel": "booking",
                "deliveryMethod": contact.get("method", "Delivery"),
                "deliveryFee": group["deliveryFee"],
                "deliveryStatus": group["deliveryStatus"],
                "status": "cancelled" if group["status"] == "cancelled" else "confirmed",
                "subtotal": subtotal,
                "total": total,
                "paid": round(paid, 2),
                "note": group["note"],
                "items": group["items"],
            }
        )
    return sales


# ────────────────────── Income & Expenses -> ledger ─────────────────────
CATEGORY_MAP = {
    "sell": "Sales",
    "inventory/ stock": "Inventory",
    "inventory/stock": "Inventory",
    "shop supplies": "Supplies",
}


def read_finance(wb) -> list[dict]:
    ws = wb["Income & Expenses"]
    entries: list[dict] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        row_date, kind, description, category, amount = row[:5]
        kind = clean(kind).lower()
        description = clean(description)
        # Columns F/G hold running totals; template rows put text there instead.
        trailing_text = [clean(cell) for cell in row[5:7] if clean(cell) and not clean(cell)[0].isdigit()]
        if kind not in ("income", "expense") or not description:
            continue
        if is_junk(description, category, *trailing_text):
            continue

        raw_category = clean(category).lower()
        mapped = CATEGORY_MAP.get(raw_category, clean(category).title() or "Other")
        # A delivery cost filed under "sell" is an expense, not revenue.
        if kind == "expense" and mapped == "Sales":
            mapped = "Delivery"

        entries.append(
            {
                "date": iso_date(row_date, "2026-09-01"),
                "kind": kind,
                "description": description,
                "category": mapped,
                "amount": num(amount),
                "method": "cash",
                "source": "import",
            }
        )
    entries.sort(key=lambda entry: entry["date"])
    return entries


# ──────────────────────────── SQL generation ────────────────────────────
def sql_str(value) -> str:
    """Quote a value for SQL, mapping blanks to NULL (for nullable columns)."""
    if value is None or value == "":
        return "null"
    return "'" + str(value).replace("'", "''") + "'"


def sql_text(value) -> str:
    """
    Quote a value for a NOT NULL text column.

    Blanks must stay as '' — sending NULL would violate the constraint and abort
    the whole seed transaction.
    """
    return "'" + ("" if value is None else str(value)).replace("'", "''") + "'"


def build_sql(stock: list[dict], sales: list[dict], finance: list[dict]) -> str:
    lines = [
        "-- Generated by tools/import_excel.py — historical data from Clothing_Shop_Management3.xlsx.",
        "-- Idempotent: safe to run more than once (natural keys sku / code are unique).",
        "",
        "begin;",
        "",
        "-- ── stock items ───────────────────────────────────────────────",
        "insert into public.stock_items (sku, name, category, size, color, cost, price, stock_in, sold)",
        "values",
    ]
    lines.append(
        ",\n".join(
            "  ({}, {}, {}, {}, {}, {}, {}, {}, {})".format(
                sql_text(item["sku"]),
                sql_text(item["name"]),
                sql_text(item["category"]),
                sql_text(item["size"]),
                sql_text(item["color"]),
                item["cost"],
                item["price"],
                item["stockIn"],
                item["sold"],
            )
            for item in stock
        )
    )
    lines += [
        "on conflict (sku) do update set",
        "  name = excluded.name, category = excluded.category, size = excluded.size,",
        "  color = excluded.color, cost = excluded.cost, price = excluded.price,",
        "  stock_in = excluded.stock_in, sold = excluded.sold;",
        "",
        "-- ── sales + line items ────────────────────────────────────────",
    ]

    for sale in sales:
        lines += [
            "with s as (",
            "  insert into public.sales (code, sold_at, customer_name, phone, address, city, channel,",
            "    delivery_method, delivery_fee, delivery_status, status, subtotal, total, paid, note)",
            "  values ({}, {}, {}, {}, {}, {}, {}, {}, {}, {}, {}, {}, {}, {}, {})".format(
                sql_text(sale["code"]),
                sql_text(sale["date"]),
                sql_text(sale["customerName"]),
                sql_text(sale["phone"]),
                sql_text(sale["address"]),
                sql_text(sale["city"]),
                sql_text(sale["channel"]),
                sql_text(sale["deliveryMethod"]),
                sale["deliveryFee"],
                sql_text(sale["deliveryStatus"]),
                sql_text(sale["status"]),
                sale["subtotal"],
                sale["total"],
                sale["paid"],
                sql_str(sale["note"]),
            ),
            "  on conflict (code) do update set customer_name = excluded.customer_name",
            "  returning id",
            ")",
            "insert into public.sale_items (sale_id, stock_item_id, name, size, color, qty, unit_price, unit_cost)",
            "select s.id, si.id, v.name, v.size, v.color, v.qty, v.unit_price, v.unit_cost",
            "from s",
            "cross join (values",
        ]
        lines.append(
            ",\n".join(
                "  ({}, {}, {}, {}, {}, {}, {})".format(
                    sql_str(item["sku"]),
                    sql_text(item["name"]),
                    sql_text(item["size"]),
                    sql_text(item["color"]),
                    item["qty"],
                    item["unitPrice"],
                    item["unitCost"],
                )
                for item in sale["items"]
            )
        )
        lines += [
            ") as v(sku, name, size, color, qty, unit_price, unit_cost)",
            "left join public.stock_items si on si.sku = v.sku",
            "where not exists (select 1 from public.sale_items existing where existing.sale_id = s.id);",
            "",
        ]

    lines += [
        "-- ── cash ledger ──────────────────────────────────────────────",
        "-- Guarded by `where not exists`: finance rows have no natural key, so this",
        "-- is what keeps a second run from doubling the ledger.",
        "insert into public.finance_entries (entry_date, kind, description, category, amount, method, source)",
        "select v.entry_date::date, v.kind, v.description, v.category, v.amount::numeric, v.method, v.source",
        "from (values",
    ]
    lines.append(
        ",\n".join(
            "  ({}, {}, {}, {}, {}, {}, {})".format(
                sql_text(entry["date"]),
                sql_text(entry["kind"]),
                sql_text(entry["description"]),
                sql_text(entry["category"]),
                entry["amount"],
                sql_text(entry["method"]),
                sql_text(entry["source"]),
            )
            for entry in finance
        )
    )
    lines += [
        ") as v(entry_date, kind, description, category, amount, method, source)",
        "where not exists (select 1 from public.finance_entries where source = 'import');",
        "",
        "commit;",
        "",
        "notify pgrst, 'reload schema';",
        "",
    ]
    return "\n".join(lines)


# ──────────────────────────────── main ──────────────────────────────────
def main() -> None:
    source = Path(sys.argv[1]).expanduser() if len(sys.argv) > 1 else None
    if not source or not source.exists():
        sys.exit("usage: import_excel.py <path to Clothing_Shop_Management3.xlsx>")

    wb = openpyxl.load_workbook(source, data_only=True)
    stock = read_stock(wb)
    stock_by_key = {match_key(item["name"], item["size"]): item for item in stock}
    delivery = read_delivery_details(wb)
    sales = read_sales(wb, stock_by_key, delivery)
    finance = read_finance(wb)

    seed = {
        "generatedAt": datetime.now().isoformat(timespec="seconds"),
        "source": source.name,
        "stockItems": stock,
        "sales": sales,
        "financeEntries": finance,
    }
    JSON_OUT.parent.mkdir(parents=True, exist_ok=True)
    JSON_OUT.write_text(json.dumps(seed, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    SQL_OUT.parent.mkdir(parents=True, exist_ok=True)
    SQL_OUT.write_text(build_sql(stock, sales, finance), encoding="utf-8")

    matched = sum(1 for sale in sales for item in sale["items"] if item["sku"])
    total_lines = sum(len(sale["items"]) for sale in sales)
    print(f"stock items      : {len(stock)}")
    print(f"sales            : {len(sales)} ({total_lines} lines, {matched} linked to a SKU)")
    print(f"ledger entries   : {len(finance)}")
    print(f"→ {JSON_OUT.relative_to(ROOT)}")
    print(f"→ {SQL_OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
